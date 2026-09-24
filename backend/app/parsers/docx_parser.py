from pathlib import Path
from uuid import UUID

import docx

from app.core.exceptions import DocumentParseError
from app.parsers.base import BaseDocumentParser, ParsedDocumentResult
from app.parsers.text_parser import chunk_text_structure_aware


class DocxDocumentParser(BaseDocumentParser):
    def parse(self, file_path: Path, document_id: UUID, filename: str) -> ParsedDocumentResult:
        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise DocumentParseError(f"Corrupted or invalid DOCX {filename}: {str(e)}") from e

        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also parse table cells if present
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_text.append(row_str)

        combined_text = "\n\n".join(paragraphs_text + table_text)
        if not combined_text.strip():
            raise DocumentParseError(f"DOCX {filename} contains no extractable text.")

        chunks = chunk_text_structure_aware(
            text=combined_text,
            document_id=document_id,
            page_number=1,  # DOCX files do not have fixed native page offsets without rendering
        )

        return ParsedDocumentResult(
            document_id=document_id,
            filename=filename,
            raw_text=combined_text,
            page_count=1,
            chunks=chunks,
            metadata={"format": "docx", "paragraph_count": len(paragraphs_text)},
        )
